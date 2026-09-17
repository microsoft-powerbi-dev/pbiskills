"""build_sample_mapping.py - regenerate examples/sample-mapping-member-loader.xlsx.

The workbook is **entirely fictional**. No real member, plan, or source system
is named and there are no data values anywhere in it - a mapping document
specifies fields, not rows. It is committed only because of that, and it is the
only mapping artifact in this repository safe to commit.

It is built to exercise the skill rather than to look tidy, so it deliberately
reproduces the things real mapping documents do:

* a title banner above the header row, so the header must be *located*
* `file name` populated once and left blank beneath, so it must be forward-filled
* two target files on one sheet, separated by a blank row
* `column name` in Excel column **C**, with `loader file column` before it
* a target size narrower than the source size (truncation risk)
* a `codes` cell that names a code set without listing its values
* a rule reading "as per current process" (underspecified - must be asked about)
* an `additional comment` that contradicts the `transformation rules` cell
* a field marked `target file needed = N`
* a COBOL-style `9(7)V99` size
* a stateful sequence rule
* a `Code Sets` tab with no recognisable header, which the reader must skip
  but still report - because that is often where the missing code values are

Run:
    python scripts/build_sample_mapping.py
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Sequence

import openpyxl

_SKILL_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = _SKILL_ROOT / "examples" / "sample-mapping-member-loader.xlsx"

TITLE = "FICTIONAL SAMPLE - Member Loader Interface Mapping - v4 draft"

HEADERS: Sequence[str] = (
    "File Name",
    "Loader File Column",
    "Column Name",
    "Base Type",
    "Size",
    "Format",
    "Mandatory",
    "Description",
    "Codes",
    "Target File Needed",
    "Additional Comment",
    "Source",
    "Source Field",
    "Source Type",
    "Source Size",
    "Source Format",
    "Source Description",
    "Transformation Rules",
    "Comments",
)

MEMBER_ROWS: List[Dict[str, Any]] = [
    {
        "File Name": "MEMBER_LOADER",
        "Loader File Column": 1,
        "Column Name": "MEMBER_ID",
        "Base Type": "CHAR",
        "Size": 12,
        "Mandatory": "Y",
        "Description": "Unique member identifier assigned at enrolment",
        "Target File Needed": "Y",
        "Source": "dbo.Member",
        "Source Field": "MemberKey",
        "Source Type": "int",
        "Source Size": 4,
        "Source Description": "Surrogate key, one row per member",
        "Transformation Rules": "Left pad with zeroes to 12",
    },
    {
        # File Name deliberately blank from here on - must be forward-filled.
        "Loader File Column": 2,
        "Column Name": "GROUP_ID",
        "Base Type": "CHAR",
        "Size": 10,
        "Mandatory": "Y",
        "Description": "Employer group the member is enrolled under",
        "Target File Needed": "Y",
        "Source": "dbo.Grp",
        "Source Field": "GroupCode",
        "Source Type": "varchar",
        "Source Size": 10,
        "Source Description": "Natural group code, unique per employer group",
        "Transformation Rules": "Direct move",
    },
    {
        "Loader File Column": 3,
        "Column Name": "LAST_NAME",
        "Base Type": "CHAR",
        "Size": 30,
        "Mandatory": "Y",
        "Description": "Member surname as held on the enrolment record",
        "Target File Needed": "Y",
        "Source": "dbo.Member",
        "Source Field": "LastName",
        "Source Type": "varchar",
        "Source Size": 40,          # wider than target - truncation risk
        "Source Description": "Free text, not validated on entry",
        "Transformation Rules": "Trim and convert to upper case",
    },
    {
        "Loader File Column": 4,
        "Column Name": "FIRST_NAME",
        "Base Type": "CHAR",
        "Size": 20,
        "Mandatory": "Y",
        "Description": "Member given name",
        "Target File Needed": "Y",
        "Source": "dbo.Member",
        "Source Field": "FirstName",
        "Source Type": "varchar",
        "Source Size": 25,          # truncation risk
        "Source Description": "Free text, not validated on entry",
        "Transformation Rules": "Trim and convert to upper case",
    },
    {
        "Loader File Column": 5,
        "Column Name": "MIDDLE_INIT",
        "Base Type": "CHAR",
        "Size": 1,
        "Mandatory": "N",
        "Description": "Middle initial",
        "Target File Needed": "Y",
        "Additional Comment": "Space if the member has no middle name",
        "Source": "dbo.Member",
        "Source Field": "MiddleName",
        "Source Type": "varchar",
        "Source Size": 25,
        "Source Description": "Full middle name where captured",
        "Transformation Rules": "First character of middle name, upper case",
    },
    {
        "Loader File Column": 6,
        "Column Name": "BIRTH_DATE",
        "Base Type": "DATE",
        "Size": 8,
        "Format": "YYYYMMDD",
        "Mandatory": "Y",
        "Description": "Member date of birth, used by the vendor for matching",
        "Target File Needed": "Y",
        "Source": "dbo.Member",
        "Source Field": "BirthDate",
        "Source Type": "date",
        "Source Size": 3,
        "Source Format": "YYYY-MM-DD",
        "Source Description": "Captured at enrolment",
        "Transformation Rules": "Format as YYYYMMDD",
    },
    {
        "Loader File Column": 7,
        "Column Name": "GENDER_CODE",
        "Base Type": "CHAR",
        "Size": 1,
        "Mandatory": "Y",
        "Description": "Gender code as required by the vendor layout",
        "Codes": "M = Male, F = Female, X = Non-binary, U = Unknown",
        "Target File Needed": "Y",
        "Additional Comment": "Default to U when the source value is null",
        "Source": "dbo.Member",
        "Source Field": "GenderCode",
        "Source Type": "char",
        "Source Size": 1,
        "Source Description": "Single character, nullable",
        "Transformation Rules": "Map per codes",
    },
    {
        "Loader File Column": 8,
        "Column Name": "RELATIONSHIP",
        "Base Type": "CHAR",
        "Size": 2,
        "Mandatory": "Y",
        "Description": "Relationship of the member to the subscriber",
        # Names a code set without listing it - a stop-and-ask.
        "Codes": "see RELATIONSHIP codes",
        "Target File Needed": "Y",
        "Source": "dbo.Enrollment",
        "Source Field": "RelationCode",
        "Source Type": "char",
        "Source Size": 2,
        "Source Description": "Internal relationship code",
        "Transformation Rules": "Translate using code table",
        "Comments": "TBD - confirm the full vendor code list before build",
    },
    {
        "Loader File Column": 9,
        "Column Name": "EFF_DATE",
        "Base Type": "DATE",
        "Size": 8,
        "Format": "YYYYMMDD",
        "Mandatory": "Y",
        "Description": "Date the member's current coverage began",
        "Target File Needed": "Y",
        "Source": "dbo.Enrollment",
        "Source Field": "EffectiveDate",
        "Source Type": "date",
        "Source Size": 3,
        "Source Format": "YYYY-MM-DD",
        "Source Description": "One row per member per plan year",
        "Transformation Rules": "Format as YYYYMMDD",
    },
    {
        "Loader File Column": 10,
        "Column Name": "TERM_DATE",
        "Base Type": "DATE",
        "Size": 8,
        "Format": "YYYYMMDD",
        "Mandatory": "N",
        "Description": "Date coverage ended; blank while coverage is active",
        "Target File Needed": "Y",
        "Additional Comment": "Spaces, not zeroes, when coverage is active",
        "Source": "dbo.Enrollment",
        "Source Field": "TerminationDate",
        "Source Type": "date",
        "Source Size": 3,
        "Source Format": "YYYY-MM-DD",
        "Source Description": "Null while the enrolment is open",
        "Transformation Rules": "Format as YYYYMMDD",
    },
    {
        "Loader File Column": 11,
        "Column Name": "PLAN_CODE",
        "Base Type": "CHAR",
        "Size": 8,
        "Mandatory": "Y",
        "Description": "Benefit plan the member is enrolled in",
        "Target File Needed": "Y",
        # Contradicts the rule cell - the exception lives here, as it does in
        # real documents.
        "Additional Comment": "Hardcode 'STANDARD' for phase 1; source the "
                              "real value from phase 2",
        "Source": "dbo.Enrollment",
        "Source Field": "PlanCode",
        "Source Type": "varchar",
        "Source Size": 8,
        "Source Description": "Benefit plan code",
        "Transformation Rules": "Direct move",
    },
    {
        "Loader File Column": 12,
        "Column Name": "PREMIUM_AMT",
        "Base Type": "NUMBER",
        "Size": "9(7)V99",
        "Format": "Implied 2 decimals, zero filled",
        "Mandatory": "N",
        "Description": "Monthly premium attributable to this member",
        "Target File Needed": "Y",
        "Source": "dbo.Billing",
        "Source Field": "PremiumAmount",
        "Source Type": "decimal",
        "Source Size": "18,2",
        "Source Description": "Latest billed premium for the member",
        "Transformation Rules": "Implied two decimals, no decimal point, "
                                "zero fill left",
    },
    {
        "Loader File Column": 13,
        "Column Name": "STATUS_FLAG",
        "Base Type": "CHAR",
        "Size": 1,
        "Mandatory": "Y",
        "Description": "Active or terminated indicator",
        "Codes": "A = Active, T = Terminated",
        "Target File Needed": "Y",
        "Source": "dbo.Enrollment",
        "Source Field": "StatusCode",
        "Source Type": "char",
        "Source Size": 1,
        "Source Description": "Internal status code",
        # Underspecified - must be asked about, never inferred.
        "Transformation Rules": "Derive from termination date as per current "
                                "process",
        "Comments": "Confirm with the enrolment team",
    },
    {
        "Loader File Column": 14,
        "Column Name": "RECORD_SEQ",
        "Base Type": "NUMBER",
        "Size": 7,
        "Mandatory": "Y",
        "Description": "Sequential record number within the transmission",
        "Target File Needed": "Y",
        "Source": "Derived",
        "Source Description": "Generated during extract",
        "Transformation Rules": "Sequence number, incrementing by 1",
        "Comments": "Confirm whether the sequence restarts per file part",
    },
    {
        "Loader File Column": 15,
        "Column Name": "LOAD_DATE",
        "Base Type": "DATE",
        "Size": 8,
        "Format": "YYYYMMDD",
        "Mandatory": "Y",
        "Description": "Date the extract was produced",
        "Target File Needed": "Y",
        "Source": "Derived",
        "Source Description": "Extract run date",
        "Transformation Rules": "Hardcode the extract run date",
    },
    {
        "Loader File Column": 16,
        "Column Name": "LEGACY_KEY",
        "Base Type": "CHAR",
        "Size": 15,
        "Mandatory": "N",
        "Description": "Legacy system member key",
        # Excluded from the output, but kept visible in the mapper.
        "Target File Needed": "N",
        "Additional Comment": "Not needed in the target file; retained here "
                              "for reconciliation against the legacy extract",
        "Source": "dbo.Member",
        "Source Field": "LegacyId",
        "Source Type": "varchar",
        "Source Size": 15,
        "Source Description": "Key from the retired platform",
        "Transformation Rules": "Direct move",
    },
]

DEPENDENT_ROWS: List[Dict[str, Any]] = [
    {
        "File Name": "DEPENDENT_LOADER",
        "Loader File Column": 1,
        "Column Name": "MEMBER_ID",
        "Base Type": "CHAR",
        "Size": 12,
        "Mandatory": "Y",
        "Description": "Subscriber the dependent is attached to",
        "Target File Needed": "Y",
        "Source": "dbo.Dependent",
        "Source Field": "MemberKey",
        "Source Type": "int",
        "Source Size": 4,
        "Source Description": "Foreign key to the subscribing member",
        "Transformation Rules": "Left pad with zeroes to 12",
    },
    {
        "Loader File Column": 2,
        "Column Name": "DEP_SEQ",
        "Base Type": "NUMBER",
        "Size": 2,
        "Mandatory": "Y",
        "Description": "Dependent sequence within the subscriber",
        "Target File Needed": "Y",
        "Source": "dbo.Dependent",
        "Source Field": "DependentSeq",
        "Source Type": "tinyint",
        "Source Size": 1,
        "Source Description": "One row per dependent per subscriber",
        "Transformation Rules": "Left pad with zeroes to 2",
    },
    {
        "Loader File Column": 3,
        "Column Name": "DEP_LAST_NAME",
        "Base Type": "CHAR",
        "Size": 30,
        "Mandatory": "Y",
        "Description": "Dependent surname",
        "Target File Needed": "Y",
        "Source": "dbo.Dependent",
        "Source Field": "LastName",
        "Source Type": "varchar",
        "Source Size": 40,
        "Source Description": "Free text",
        "Transformation Rules": "Trim and convert to upper case",
    },
    {
        "Loader File Column": 4,
        "Column Name": "DEP_BIRTH_DATE",
        "Base Type": "DATE",
        "Size": 8,
        "Format": "YYYYMMDD",
        "Mandatory": "Y",
        "Description": "Dependent date of birth",
        "Target File Needed": "Y",
        "Source": "dbo.Dependent",
        "Source Field": "BirthDate",
        "Source Type": "date",
        "Source Size": 3,
        "Source Format": "YYYY-MM-DD",
        "Source Description": "Captured at enrolment",
        "Transformation Rules": "Format as YYYYMMDD",
    },
]

#: A tab with no recognisable header row. The reader must skip it and still
#: report it, because a named-but-unlisted code set is usually hiding on one.
CODE_SET_NOTES: Sequence[Sequence[str]] = (
    ("Code set reference - maintained separately by the interface team",),
    ("",),
    ("RELATIONSHIP", "pending vendor confirmation"),
    ("GENDER", "see the Member Loader tab"),
)


def _write_mapping_sheet(wb: openpyxl.Workbook, name: str,
                         blocks: Sequence[Sequence[Dict[str, Any]]]) -> None:
    """A mapping sheet: title banner, header row, then blocks split by a blank row."""
    ws = wb.create_sheet(name)
    ws.append([TITLE])
    ws.append(list(HEADERS))
    for index, block in enumerate(blocks):
        if index:
            ws.append([])
        for row in block:
            ws.append([row.get(header) for header in HEADERS])


def build() -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    _write_mapping_sheet(wb, "Member Loader", (MEMBER_ROWS, DEPENDENT_ROWS))

    notes = wb.create_sheet("Code Sets")
    for row in CODE_SET_NOTES:
        notes.append(list(row))

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print("wrote", build())
