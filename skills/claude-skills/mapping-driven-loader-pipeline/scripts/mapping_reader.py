"""mapping_reader.py - mapping workbook to mapping_digest.json.

Deterministic, openpyxl-only, and deliberately free of interpretation. It
locates the header row, normalises the 19 known column names against their
aliases, forward-fills merged/blank cells, groups rows into target files, and
records every column and sheet it could not make sense of.

The digest is the seam, the same way `sql-server-schema`'s `schema_digest.json`
is: it is reviewable before anything is generated, it makes the downstream
mapper and code generation testable against a fixture, and two digests of
successive spec revisions diff cleanly so a mapping-document revision produces
a reviewable change rather than an archaeology exercise.

Parse and validate are strictly separate, following
`extract-engine/scripts/extract_engine/config_loader.py`: reading never raises
on a content problem, it accumulates `Issue` records into the digest. The one
exception is an ambiguous target-name column, where continuing would silently
misalign every downstream artifact - see `_resolve_column_name_source`.

Run:
    python mapping_reader.py read   --workbook <path> [--out digest.json]
    python mapping_reader.py report --workbook <path>
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

SCHEMA_VERSION = "1.0"
GENERATOR_VERSION = "1.0"

#: How far down a sheet to look for the header row. Mapping documents carry
#: title banners, revision stamps and merged section headers above the real
#: header, but never twenty rows of them.
HEADER_SEARCH_DEPTH = 20

#: Minimum known-column matches for a row to be accepted as the header. Four
#: is low enough to tolerate a document that renames half its columns and high
#: enough that a notes row full of prose never wins.
HEADER_MIN_MATCHES = 4

#: Excel column C, zero-based. The documented position of the target column
#: name; used only as a fallback when no header alias matches.
COLUMN_NAME_FALLBACK_INDEX = 2

#: Canonical field -> accepted header aliases. The canonical name is listed
#: first so it is always its own alias. Order within a tuple is irrelevant;
#: matching is exact against the normalised form.
ALIASES: Dict[str, Tuple[str, ...]] = {
    "file_name": (
        "file name", "file", "filename", "target file", "output file",
        "loader file name", "target file name", "extract file",
    ),
    "column_name": (
        "column name", "target column", "field name", "column", "target field",
        "target column name", "field", "column nm",
    ),
    "loader_file_column": (
        "loader file column", "loader column", "loader field", "position",
        "seq", "seq no", "ordinal", "column no", "column number", "field no",
        "loader position", "order",
    ),
    "base_type": (
        "base type", "type", "target type", "data type", "datatype",
        "base data type", "target data type",
    ),
    "size": (
        "size", "length", "len", "target size", "width", "field size",
        "target length", "field length",
    ),
    "format": (
        "format", "target format", "output format", "mask", "picture",
        "target mask", "edit mask",
    ),
    "mandatory": (
        "mandatory", "required", "req", "mandatory optional", "m o",
        "mandatory y n", "is mandatory", "required y n",
    ),
    "description": (
        "description", "desc", "field description", "business description",
        "definition", "target description", "column description",
    ),
    "codes": (
        "codes", "code", "valid values", "values", "domain", "code values",
        "allowed values", "code set", "valid codes", "lookup values",
    ),
    "target_file_needed": (
        "target file needed", "needed", "required in target", "include",
        "in scope", "send", "needed in target file", "target needed",
        "include in file",
    ),
    "additional_comment": (
        "additional comment", "additional comments", "notes", "remarks",
        "comment 2", "additional notes", "further comments",
    ),
    "source": (
        "source", "source table", "source system", "src", "source object",
        "from", "source entity", "src table",
    ),
    "source_field": (
        "source field", "source column", "src field", "source attribute",
        "src column", "source fields", "source column name",
    ),
    "source_type": (
        "source type", "src type", "source data type", "source datatype",
        "src data type",
    ),
    "source_size": (
        "source size", "src size", "source length", "src length",
        "source width",
    ),
    "source_format": (
        "source format", "src format", "source mask", "src mask",
        "source picture",
    ),
    "source_description": (
        "source description", "src description", "source definition",
        "source desc", "source business description",
    ),
    "transformation_rules": (
        "transformation rules", "transformation rule", "transformation",
        "rule", "logic", "mapping rule", "transformation logic",
        "business rule", "transform rule", "transformation logic rule",
        "rules", "derivation",
    ),
    "comments": (
        "comments", "comment", "open items", "questions", "issues",
        "open questions", "review comments",
    ),
}

#: Aliases whose sense is inverted relative to `mandatory`. A column headed
#: "Nullable" means the opposite of one headed "Mandatory", and reading it the
#: wrong way round inverts the null policy on every field in the document.
INVERTED_MANDATORY_ALIASES: Tuple[str, ...] = (
    "nullable", "null", "nulls allowed", "optional", "allow null",
)

#: Columns forward-filled when blank. A mapping sheet names the file once and
#: leaves the cell blank for the rows beneath it; the same happens to `source`
#: where one table serves a run of fields. A blank here means "same as above",
#: never "none".
FORWARD_FILL_FIELDS: Tuple[str, ...] = ("file_name", "source")

TRUE_TOKENS = frozenset((
    "y", "yes", "true", "1", "m", "mandatory", "required", "not null",
    "notnull", "x", "needed", "include",
))
FALSE_TOKENS = frozenset((
    "n", "no", "false", "0", "o", "optional", "null", "nullable",
    "not needed", "not required", "exclude", "na", "n a",
))

#: Values in `source` that mean "there is no source object", so a blank
#: `source_field` on such a row is expected rather than an omission.
NO_SOURCE_TOKENS = frozenset((
    "derived", "n a", "na", "none", "hardcoded", "hard coded", "constant",
    "default", "literal", "system generated", "generated",
))

_PASSTHROUGH_RULES = frozenset((
    "direct", "direct move", "as is", "asis", "1 1", "straight move",
    "direct map", "no transformation", "same", "move", "direct mapping",
))

#: Rule text that is *entirely* a literal. Checked as a whole-value match, not
#: as a substring: "zeroes" alone is a constant, but "left pad with zeroes to
#: 12" is a pad.
_BARE_LITERAL_RULES = frozenset((
    "spaces", "space", "blank", "blanks", "zeroes", "zeros", "zero",
    "all spaces", "all zeroes", "all zeros", "low values", "high values",
    "null", "empty",
))

#: Rule text that states no derivation. These are the rows that produce a wrong
#: loader file when inferred instead of asked about.
_UNDERSPECIFIED_PATTERNS: Tuple[str, ...] = (
    "per business logic", "as per current process", "same as legacy",
    "as per existing", "business logic", "tbd", "to be defined",
    "to be confirmed", "as discussed", "per current", "refer to legacy",
)

#: Text in `comments` / `additional_comment` that marks an unresolved item.
_OPEN_QUESTION_PATTERNS: Tuple[str, ...] = (
    "tbd", "to be confirmed", "to be decided", "confirm", "check with",
    "pending", "clarify", "question", "unclear", "not sure", "???",
)

_SIZE_PATTERN = re.compile(r"^(\d+)\s*(?:[,.]\s*(\d+))?$")
_PIC_PATTERN = re.compile(r"9\s*\(\s*(\d+)\s*\)\s*(?:v\s*9{1,}|\.\s*9{1,})?",
                          re.IGNORECASE)
_TYPE_SIZE_PATTERN = re.compile(r"\(\s*(\d+)\s*(?:,\s*(\d+)\s*)?\)")


# ---------------------------------------------------------------------------
# Normalisation helpers
# ---------------------------------------------------------------------------

def normalize_header(value: Any) -> str:
    """A header cell reduced to its comparable form.

    Lowercased, punctuation dropped, whitespace collapsed. This is what makes
    "Source Field", "source_field" and "Source field:" the same column, which
    they always are in a hand-maintained document.
    """
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: str) -> str:
    """A file name turned into a directory-safe identifier."""
    text = re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower())
    return text.strip("_") or "unnamed"


def _cell(value: Any) -> Optional[str]:
    """A data cell as trimmed text, or None when empty.

    Excel hands back floats for integer-looking cells, so 1.0 becomes "1"
    rather than "1.0" - an ordinal of "1.0" sorts and prints wrongly
    everywhere downstream.
    """
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def _boolean(value: Any, *, inverted: bool = False) -> Optional[bool]:
    """A Y/N-ish cell as a boolean, or None when it says nothing recognisable.

    ``inverted`` flips the sense for a column headed "Nullable" or "Optional",
    which means the opposite of one headed "Mandatory". Reading that the wrong
    way round inverts the null policy on every field in the document, so the
    sense is carried explicitly rather than assumed.
    """
    text = normalize_header(value)
    if not text:
        return None
    first = text.split(" ")[0]
    for token in (text, first):
        if token in TRUE_TOKENS:
            return not inverted
        if token in FALSE_TOKENS:
            return inverted
    return None


def parse_size(size_text: Optional[str],
               base_type: Optional[str] = None) -> Tuple[Optional[int], Optional[int]]:
    """Width and scale from a `size` cell, and from a type like `DECIMAL(9,2)`.

    Handles the three spellings these documents use interchangeably: a plain
    number, a `precision,scale` pair, and COBOL-style `9(7)V99`. Returns
    (None, None) rather than raising when it cannot tell - an unparseable size
    becomes an Issue, not an exception.

    For a PIC clause the returned width is the **emitted character width**,
    integer digits plus implied decimals: `9(7)V99` is nine characters on the
    line, not seven, because `V` is an implied point that occupies no
    character. Getting this wrong under-sizes the field check by the scale.
    """
    text = (size_text or "").strip()
    if text:
        match = _SIZE_PATTERN.match(text)
        if match:
            return int(match.group(1)), int(match.group(2)) if match.group(2) else None
        pic = _PIC_PATTERN.search(text)
        if pic:
            lowered = text.lower()
            scale = (lowered.count("9", lowered.find("v"))
                     if "v" in lowered else None)
            return int(pic.group(1)) + (scale or 0), scale
    if base_type:
        embedded = _TYPE_SIZE_PATTERN.search(base_type)
        if embedded:
            return (int(embedded.group(1)),
                    int(embedded.group(2)) if embedded.group(2) else None)
    return None, None


def classify_rule(rule_text: Optional[str]) -> str:
    """The rule's shape, per references/mapping-document-schema.md.

    A hint for the generator and a column in the mapper's rule catalog - never
    a substitute for reading the rule. `underspecified` and `unclassified` both
    mean "ask", and are reported as Issues.

    The checks are ordered by specificity, and the order is load-bearing.
    "Hardcode the extract run date" mentions a date but is a constant, so the
    constant test precedes the date test. "Left pad with zeroes to 12"
    mentions zeroes but is a pad, so the bare-literal tokens (`spaces`,
    `zeroes`) only count as a constant when they are the *whole* rule. And
    "Format as YYYYMMDD" contains "format", so the date test precedes the
    string test.
    """
    text = normalize_header(rule_text)
    if not text:
        return "passthrough"
    if text in _PASSTHROUGH_RULES:
        return "passthrough"
    for pattern in _UNDERSPECIFIED_PATTERNS:
        if pattern in text:
            return "underspecified"
    if any(k in text for k in ("hardcode", "hard code", "constant", "literal",
                               "default to", "defaulted to", "set to",
                               "always ")):
        return "constant"
    if text in _BARE_LITERAL_RULES:
        return "constant"
    if any(k in text for k in ("sequence", "record number", "incrementing",
                               "row number", "running number")):
        return "row_sequence"
    if any(k in text for k in ("sum", "count of", "max of", "min of", "average",
                               "latest", "most recent", "earliest", "total of")):
        return "aggregate"
    if any(k in text for k in ("if ", "when ", "case ", "else", "otherwise",
                               "depending on")):
        return "conditional"
    if any(k in text for k in ("code", "translate", "map per", "lookup",
                               "cross reference", "xref")):
        return "code_lookup"
    if any(k in text for k in ("concat", "append", "combine", " plus ",
                               "followed by")):
        return "concatenation"
    if any(k in text for k in ("implied", "decimal", "zero fill", "zero filled",
                               "sign", "no decimal point")):
        return "numeric_render"
    if any(k in text for k in ("date", "yyyy", "ccyy", "mmdd", "timestamp")):
        return "date_render"
    if any(k in text for k in ("trim", "upper", "lower", "pad", "substring",
                               "left", "right", "strip", "truncate", "format",
                               "remove", "replace", "first character")):
        return "string_transform"
    return "unclassified"


def _looks_open(text: Optional[str]) -> bool:
    normalized = normalize_header(text)
    if not normalized:
        return False
    return any(p in normalized for p in _OPEN_QUESTION_PATTERNS) or "?" in str(text)


# ---------------------------------------------------------------------------
# Digest model
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Issue:
    severity: str            # "error" | "warning" | "info"
    message: str
    sheet: Optional[str] = None
    row: Optional[int] = None
    column: Optional[str] = None
    file_name: Optional[str] = None
    field: Optional[str] = None

    def __str__(self) -> str:
        where = ":".join(str(p) for p in (self.sheet, self.row, self.column) if p)
        return "[{}] {}{}".format(self.severity, where + " - " if where else "",
                                  self.message)


@dataclass
class SheetRead:
    name: str
    header_row: int
    field_rows: int
    unmatched_columns: List[str] = field(default_factory=list)
    column_name_source: str = "header"
    mandatory_sense: str = "mandatory"


class MappingReadError(RuntimeError):
    """The document's shape is not what every downstream artifact assumes.

    Raised only where continuing would silently misalign output rather than
    produce a reportable problem.
    """


# ---------------------------------------------------------------------------
# Header resolution
# ---------------------------------------------------------------------------

def _alias_index() -> Dict[str, str]:
    index: Dict[str, str] = {}
    for canonical, aliases in ALIASES.items():
        for alias in aliases:
            index[normalize_header(alias)] = canonical
    for alias in INVERTED_MANDATORY_ALIASES:
        index[normalize_header(alias)] = "mandatory"
    return index


_ALIAS_INDEX = _alias_index()


def locate_header_row(rows: Sequence[Sequence[Any]]) -> Tuple[int, int]:
    """The index of the header row and how many known columns it matched.

    Scored rather than assumed: these documents put titles, revision stamps
    and merged section labels above the header, and the header is simply the
    row that looks most like one.
    """
    best_index, best_score = -1, 0
    for index, row in enumerate(rows[:HEADER_SEARCH_DEPTH]):
        score = sum(1 for cell in row
                    if normalize_header(cell) in _ALIAS_INDEX)
        if score > best_score:
            best_index, best_score = index, score
    return best_index, best_score


def map_columns(header: Sequence[Any]) -> Tuple[Dict[str, int], List[str], str]:
    """Canonical field -> column index, plus unmatched headers and the sense
    the `mandatory` column was read in.

    First match wins: a document with both "Description" and "Target
    Description" keeps the leftmost, which is the one the field table is built
    around.
    """
    mapping: Dict[str, int] = {}
    unmatched: List[str] = []
    mandatory_sense = "mandatory"
    for index, cell in enumerate(header):
        normalized = normalize_header(cell)
        if not normalized:
            continue
        canonical = _ALIAS_INDEX.get(normalized)
        if canonical is None:
            unmatched.append(str(cell).strip())
        elif canonical not in mapping:
            mapping[canonical] = index
            if canonical == "mandatory" and normalized in {
                    normalize_header(a) for a in INVERTED_MANDATORY_ALIASES}:
                mandatory_sense = "nullable"
    return mapping, unmatched, mandatory_sense


def _resolve_column_name_source(mapping: Dict[str, int],
                                header: Sequence[Any]) -> str:
    """Whether the target name came from a header match or from column C.

    Raises when both resolve and disagree. That disagreement means the
    document's shape differs from the expected shape, and every ordinal,
    transform name and layout entry downstream would be built on the wrong
    column - a failure that surfaces as bad data in a loader, days later.
    """
    header_index = mapping.get("column_name")
    has_fallback = (len(header) > COLUMN_NAME_FALLBACK_INDEX
                    and _cell(header[COLUMN_NAME_FALLBACK_INDEX]) is not None)
    if header_index is None:
        if not has_fallback:
            raise MappingReadError(
                "no column matched a `column name` alias and Excel column C is "
                "empty; cannot identify the target column name")
        mapping["column_name"] = COLUMN_NAME_FALLBACK_INDEX
        return "position_c"
    if header_index != COLUMN_NAME_FALLBACK_INDEX and has_fallback:
        raise MappingReadError(
            "ambiguous target column name: header {!r} is at column {} but "
            "column C holds {!r}. Ask which column is the target column name "
            "before generating anything.".format(
                str(header[header_index]).strip(), header_index + 1,
                str(header[COLUMN_NAME_FALLBACK_INDEX]).strip()))
    return "header"


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

def _read_sheet(ws: Any, issues: List[Issue]) -> Tuple[Optional[SheetRead],
                                                       List[Dict[str, Any]]]:
    """One worksheet's field rows as raw canonical dicts."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return None, []

    header_index, score = locate_header_row(rows)
    if header_index < 0 or score < HEADER_MIN_MATCHES:
        issues.append(Issue("info", "no header row matched (best score {})"
                            .format(score), sheet=ws.title))
        return None, []

    header = rows[header_index]
    mapping, unmatched, mandatory_sense = map_columns(header)
    column_name_source = _resolve_column_name_source(mapping, header)

    for canonical in ("transformation_rules", "source_field", "base_type", "size"):
        if canonical not in mapping:
            issues.append(Issue(
                "warning", "no `{}` column found on this sheet".format(
                    canonical.replace("_", " ")), sheet=ws.title))

    read = SheetRead(name=ws.title, header_row=header_index + 1, field_rows=0,
                     unmatched_columns=unmatched,
                     column_name_source=column_name_source,
                     mandatory_sense=mandatory_sense)

    def value_of(row: Sequence[Any], canonical: str) -> Optional[str]:
        index = mapping.get(canonical)
        if index is None or index >= len(row):
            return None
        return _cell(row[index])

    inverted = mandatory_sense == "nullable"
    carried: Dict[str, Optional[str]] = {k: None for k in FORWARD_FILL_FIELDS}
    out: List[Dict[str, Any]] = []

    for offset, row in enumerate(rows[header_index + 1:]):
        row_number = header_index + 2 + offset
        if all(cell is None for cell in row):
            continue

        record: Dict[str, Any] = {canonical: value_of(row, canonical)
                                  for canonical in ALIASES}
        for key in FORWARD_FILL_FIELDS:
            if record.get(key):
                carried[key] = record[key]
            else:
                record[key] = carried[key]

        if not record.get("column_name"):
            # A row carrying a file name but no target column is a section
            # separator, not a field.
            continue

        record["mandatory"] = _boolean(
            row[mapping["mandatory"]] if "mandatory" in mapping
            and mapping["mandatory"] < len(row) else None, inverted=inverted)
        needed = _boolean(
            row[mapping["target_file_needed"]] if "target_file_needed" in mapping
            and mapping["target_file_needed"] < len(row) else None)
        record["target_file_needed"] = True if needed is None else needed
        record["sheet"] = ws.title
        record["row"] = row_number
        out.append(record)

    read.field_rows = len(out)
    return read, out


def _finalize_field(record: Dict[str, Any], ordinal: int,
                    issues: List[Issue]) -> Dict[str, Any]:
    """One raw row turned into a digest field entry, with its own findings."""
    file_name = record.get("file_name") or "UNNAMED"
    name = record["column_name"]
    where = {"sheet": record.get("sheet"), "row": record.get("row"),
             "file_name": file_name, "field": name}

    size, scale = parse_size(record.get("size"), record.get("base_type"))
    if record.get("size") and size is None:
        issues.append(Issue("warning", "unparseable size {!r}".format(
            record["size"]), column="size", **where))

    source_size, _ = parse_size(record.get("source_size"),
                                record.get("source_type"))
    if size is not None and source_size is not None and source_size > size:
        issues.append(Issue(
            "warning", "truncation risk: source size {} exceeds target size {}"
            .format(source_size, size), column="size", **where))

    classification = classify_rule(record.get("transformation_rules"))
    if classification == "underspecified":
        issues.append(Issue(
            "error", "underspecified rule: {!r}".format(
                record.get("transformation_rules")),
            column="transformation rules", **where))
    elif classification == "unclassified":
        issues.append(Issue(
            "warning", "rule could not be classified: {!r}".format(
                record.get("transformation_rules")),
            column="transformation rules", **where))
    elif classification == "aggregate":
        issues.append(Issue(
            "warning", "rule implies an aggregate, which changes the row "
                       "grain and belongs in the source query",
            column="transformation rules", **where))
    elif classification == "row_sequence":
        issues.append(Issue(
            "warning", "stateful sequence rule: confirm per-file or per-part",
            column="transformation rules", **where))

    if not record.get("transformation_rules"):
        issues.append(Issue(
            "info", "no transformation rule stated; read as a direct move - "
                    "confirm this document's convention once per file",
            column="transformation rules", **where))

    codes = record.get("codes")
    if codes and not re.search(r"[=:\-]", codes) and len(codes.split()) <= 6:
        issues.append(Issue(
            "error", "codes names a code set without listing values: {!r}"
            .format(codes), column="codes", **where))

    source = record.get("source")
    source_normalized = normalize_header(source)
    has_source_object = bool(source) and source_normalized not in NO_SOURCE_TOKENS
    if not record.get("source_field") and has_source_object \
            and classification != "constant":
        issues.append(Issue(
            "error", "no source field for a non-constant field",
            column="source field", **where))
    if not source and classification != "constant":
        issues.append(Issue(
            "error", "no source stated for a non-constant field",
            column="source", **where))

    if record.get("mandatory") and classification == "passthrough" \
            and not record.get("additional_comment"):
        issues.append(Issue(
            "info", "mandatory field with a direct move: confirm the source is "
                    "non-nullable or state a default",
            column="mandatory", **where))

    for column in ("comments", "additional_comment"):
        if _looks_open(record.get(column)):
            issues.append(Issue(
                "warning", "open question in {}: {!r}".format(
                    column.replace("_", " "), record[column]),
                column=column.replace("_", " "), **where))

    source_fields = [p.strip() for p in re.split(r"[,+;/]", record["source_field"])
                     if p.strip()] if record.get("source_field") else []

    return {
        "ordinal": ordinal,
        "column_name": name,
        "loader_file_column": record.get("loader_file_column"),
        "base_type": record.get("base_type"),
        "size": size,
        "scale": scale,
        "size_raw": record.get("size"),
        "format": record.get("format"),
        "mandatory": record.get("mandatory"),
        "description": record.get("description"),
        "codes": codes,
        "target_file_needed": record.get("target_file_needed", True),
        "additional_comment": record.get("additional_comment"),
        "source": source,
        "source_field": record.get("source_field"),
        "source_fields": source_fields,
        "source_type": record.get("source_type"),
        "source_size": record.get("source_size"),
        "source_format": record.get("source_format"),
        "source_description": record.get("source_description"),
        "transformation_rules": record.get("transformation_rules"),
        "rule_classification": classification,
        "comments": record.get("comments"),
        "sheet": record.get("sheet"),
        "row": record.get("row"),
    }


def _order_fields(records: Sequence[Dict[str, Any]], file_name: str,
                  issues: List[Issue]) -> List[Dict[str, Any]]:
    """Field rows in loader order, with ordinal problems reported.

    A numeric `loader file column` is authoritative; gaps and duplicates in it
    shift every following field in the delimited line, so they are errors
    rather than notes. Without that column, sheet row order is the order - and
    the digest says so.
    """
    numeric: List[Tuple[int, Dict[str, Any]]] = []
    for record in records:
        raw = record.get("loader_file_column")
        try:
            numeric.append((int(str(raw).strip()), record))
        except (TypeError, ValueError):
            numeric = []
            break

    if numeric:
        numeric.sort(key=lambda pair: pair[0])
        declared = [ordinal for ordinal, _ in numeric]
        duplicates = sorted({o for o in declared if declared.count(o) > 1})
        if duplicates:
            issues.append(Issue(
                "error", "duplicate loader ordinals {}".format(duplicates),
                column="loader file column", file_name=file_name))
        expected = list(range(declared[0], declared[0] + len(declared)))
        if declared != expected:
            issues.append(Issue(
                "error", "loader ordinals are not contiguous: {}".format(declared),
                column="loader file column", file_name=file_name))
        ordered = [record for _, record in numeric]
    else:
        issues.append(Issue(
            "warning", "no usable numeric `loader file column`; using sheet row "
                       "order as the field order", file_name=file_name))
        ordered = sorted(records, key=lambda r: (r.get("sheet") or "",
                                                 r.get("row") or 0))

    names = [r["column_name"] for r in ordered]
    for name in sorted({n for n in names if names.count(n) > 1}):
        issues.append(Issue(
            "error", "duplicate target column name {!r}".format(name),
            file_name=file_name, field=name))

    return ordered


def read_workbook(path: str) -> Dict[str, Any]:
    """Parse the mapping workbook into the digest. No interpretation.

    openpyxl is imported lazily so `--help` and the unit tests that exercise
    the pure helpers work on a machine without it.
    """
    import openpyxl

    workbook_path = Path(path)
    sha256 = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(workbook_path, data_only=True, read_only=True)

    issues: List[Issue] = []
    sheets_read: List[SheetRead] = []
    sheets_skipped: List[Dict[str, str]] = []
    records: List[Dict[str, Any]] = []

    for sheet_name in wb.sheetnames:
        read, rows = _read_sheet(wb[sheet_name], issues)
        if read is None:
            sheets_skipped.append({"sheet": sheet_name,
                                   "reason": "no header row matched"})
            continue
        sheets_read.append(read)
        records.extend(rows)

    if not records:
        issues.append(Issue("error", "no field rows found in any sheet"))

    grouped: Dict[str, List[Dict[str, Any]]] = {}
    for record in records:
        grouped.setdefault(record.get("file_name") or "UNNAMED", []).append(record)

    files: List[Dict[str, Any]] = []
    for file_name in sorted(grouped):
        ordered = _order_fields(grouped[file_name], file_name, issues)
        fields = [_finalize_field(record, index, issues)
                  for index, record in enumerate(ordered, start=1)]
        included = [f for f in fields if f["target_file_needed"]]
        if not included:
            issues.append(Issue(
                "warning", "every field is marked not needed; the whole file "
                           "may be out of scope", file_name=file_name))
        files.append({
            "file_name": file_name,
            "slug": slugify(file_name),
            "field_count": len(fields),
            "included_count": len(included),
            "excluded": [f["column_name"] for f in fields
                         if not f["target_file_needed"]],
            "source_objects": sorted({f["source"] for f in fields
                                      if f["source"]}),
            "rule_classifications": _count_by(fields, "rule_classification"),
            "fields": fields,
        })

    unmatched = sorted({c for read in sheets_read for c in read.unmatched_columns})
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_by": "mapping_reader.py {}".format(GENERATOR_VERSION),
        "source_workbook": workbook_path.name,
        "source_sha256": sha256,
        "sheets_read": {r.name: r.field_rows for r in sheets_read},
        "sheets_skipped": sheets_skipped,
        "unmatched_columns": unmatched,
        "column_name_source": sheets_read[0].column_name_source
                              if sheets_read else None,
        "mandatory_sense": sheets_read[0].mandatory_sense
                           if sheets_read else None,
        "files": files,
        "issues": [asdict(i) for i in issues],
        "issue_counts": _count_severities(issues),
    }


def _count_by(items: Sequence[Dict[str, Any]], key: str) -> Dict[str, int]:
    counts: Dict[str, int] = {}
    for item in items:
        counts[item[key]] = counts.get(item[key], 0) + 1
    return dict(sorted(counts.items()))


def _count_severities(issues: Sequence[Issue]) -> Dict[str, int]:
    counts = {"error": 0, "warning": 0, "info": 0}
    for issue in issues:
        counts[issue.severity] = counts.get(issue.severity, 0) + 1
    return counts


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _print(payload: Any) -> None:
    print(json.dumps(payload, indent=2, sort_keys=True, default=str))


def cmd_read(args: argparse.Namespace) -> int:
    digest = read_workbook(args.workbook)
    if args.out:
        out_path = Path(args.out)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            json.dumps(digest, indent=2, sort_keys=True, default=str) + "\n",
            encoding="utf-8", newline="\n")
        _print({"ok": digest["issue_counts"]["error"] == 0,
                "out": str(out_path),
                "files": [f["file_name"] for f in digest["files"]],
                "issue_counts": digest["issue_counts"]})
    else:
        _print(digest)
    return 0 if digest["issue_counts"]["error"] == 0 else 1


def cmd_report(args: argparse.Namespace) -> int:
    """A human summary. Read this before writing the transformation mapper."""
    digest = read_workbook(args.workbook)
    lines = ["Mapping document: {}".format(digest["source_workbook"]),
             "sha256: {}".format(digest["source_sha256"][:16]),
             "target name resolved by: {}".format(digest["column_name_source"]),
             "mandatory column read as: {}".format(digest["mandatory_sense"]),
             ""]
    for entry in digest["files"]:
        lines.append("{}  ({} fields, {} included, {} excluded)".format(
            entry["file_name"], entry["field_count"], entry["included_count"],
            len(entry["excluded"])))
        lines.append("  sources: {}".format(", ".join(entry["source_objects"])
                                            or "none stated"))
        lines.append("  rules:   {}".format(", ".join(
            "{}={}".format(k, v)
            for k, v in entry["rule_classifications"].items())))
        lines.append("")
    if digest["sheets_skipped"]:
        lines.append("Sheets skipped (check these for code-set values):")
        for skipped in digest["sheets_skipped"]:
            lines.append("  {} - {}".format(skipped["sheet"], skipped["reason"]))
        lines.append("")
    if digest["unmatched_columns"]:
        lines.append("Unmatched columns: {}".format(
            ", ".join(digest["unmatched_columns"])))
        lines.append("")
    counts = digest["issue_counts"]
    lines.append("Issues: {} error, {} warning, {} info".format(
        counts["error"], counts["warning"], counts["info"]))
    for raw in digest["issues"]:
        if raw["severity"] != "info":
            lines.append("  " + str(Issue(**{k: v for k, v in raw.items()})))
    print("\n".join(lines))
    return 0 if counts["error"] == 0 else 1


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="mapping_reader.py",
        description="Read a mapping document into a deterministic digest.")
    sub = parser.add_subparsers(dest="command", required=True)

    read = sub.add_parser("read", help="emit mapping_digest.json")
    read.add_argument("--workbook", required=True, help="path to the mapping document")
    read.add_argument("--out", help="write the digest here instead of stdout")
    read.set_defaults(func=cmd_read)

    report = sub.add_parser("report", help="human-readable summary")
    report.add_argument("--workbook", required=True)
    report.set_defaults(func=cmd_report)

    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        return int(args.func(args))
    except MappingReadError as exc:
        print("mapping document shape is ambiguous: {}".format(exc),
              file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
